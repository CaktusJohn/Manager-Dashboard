"""Однократная загрузка исходных Excel-файлов в SQLite."""

from datetime import date, datetime
import sqlite3
from typing import BinaryIO

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from database import get_connection, get_meta


IMPORT_FLAG = "initial_import_completed"
SOURCES = {
    "blanks": ("record_id", "date", "batch_id", "employee", "plan_qty", "fact_qty"),
    "welding_needs": ("record_id", "date", "batch_id", "need_qty"),
    "documentation": ("record_id", "date", "employee", "plan_docs", "fact_docs"),
}


def initial_import_completed() -> bool:
    return get_meta(IMPORT_FLAG) == "yes"


def to_iso_date(value: object) -> str | None:
    """Приводит дату Excel к формату YYYY-MM-DD, сохраняя пустое значение."""
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        converted = from_excel(value)
        return converted.date().isoformat() if isinstance(converted, datetime) else converted.isoformat()
    return str(value)


def read_source(
    file: BinaryIO, source_name: str, expected_headers: tuple[str, ...]
) -> list[tuple[object, ...]]:
    """Читает лист «Данные» и сверяет его заголовки с ожидаемой схемой."""
    workbook = load_workbook(file, data_only=True, read_only=True)
    try:
        if "Данные" not in workbook.sheetnames:
            raise ValueError(f"В файле «{source_name}» нет листа «Данные».")
        sheet = workbook["Данные"]
        headers = tuple(cell.value for cell in next(sheet.iter_rows(max_row=1)))
        if headers != expected_headers:
            raise ValueError(f"Неверные заголовки в файле «{source_name}».")

        result = []
        for source_row, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(value is not None for value in row):
                continue
            values = list(row[: len(expected_headers)])
            values[1] = to_iso_date(values[1])
            result.append((*values, source_row))
        return result
    finally:
        workbook.close()


def import_initial_files(files: dict[str, BinaryIO]) -> dict[str, int]:
    """Загружает три выбранных файла одной транзакцией и блокирует повторный импорт."""
    if initial_import_completed():
        raise ValueError("Первоначальный импорт уже выполнен.")

    if set(files) != set(SOURCES) or any(file is None for file in files.values()):
        raise ValueError("Выберите все три исходных файла.")

    loaded: list[tuple[str, tuple[str, ...], list[tuple[object, ...]]]] = []
    for table, columns in SOURCES.items():
        file = files[table]
        loaded.append((table, columns, read_source(file, file.name, columns)))

    counts: dict[str, int] = {}
    with get_connection() as connection:
        try:
            for table, columns, rows in loaded:
                names = (*columns, "source_row")
                placeholders = ", ".join("?" for _ in names)
                connection.executemany(
                    f"INSERT INTO {table} ({', '.join(names)}) VALUES ({placeholders})",
                    rows,
                )
                counts[table] = len(rows)
            connection.execute(
                "INSERT INTO app_meta (key, value) VALUES (?, ?)",
                (IMPORT_FLAG, "yes"),
            )
        except sqlite3.Error:
            connection.rollback()
            raise
    return counts
